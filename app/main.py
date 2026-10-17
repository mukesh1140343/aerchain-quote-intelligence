"""Aerchain Quote Intelligence API: FastAPI + Groq."""
import csv, io, json, os
from pathlib import Path
from typing import Any
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field
from openai import OpenAI
from openpyxl import load_workbook
from docx import Document
from pypdf import PdfReader

APP_DIR = Path(__file__).resolve().parent
DATA_DIR = APP_DIR.parent / "data"
MODEL = os.getenv("GROQ_MODEL", "openai/gpt-oss-120b")
FRONTEND_ORIGIN = "https://magenta-cat-044395.netlify.app"
app = FastAPI(title="Aerchain Quote Intelligence API")
app.add_middleware(
    CORSMiddleware,
    allow_origins=[FRONTEND_ORIGIN, "http://localhost:5173", "http://127.0.0.1:5173", "null"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

class AnalystRequest(BaseModel):
    question: str
    comparison_data: list[dict[str, Any]] = Field(default_factory=list)
    assumptions: dict[str, Any] = Field(default_factory=dict)

def client() -> OpenAI:
    key = os.getenv("GROQ_API_KEY")
    if not key:
        raise HTTPException(503, "AI is not configured. Add GROQ_API_KEY in Render environment variables.")
    return OpenAI(api_key=key, base_url="https://api.groq.com/openai/v1")

def ask_model(system: str, prompt: str, max_tokens: int = 1400) -> str:
    try:
        response = client().chat.completions.create(
            model=MODEL,
            messages=[{"role": "system", "content": system}, {"role": "user", "content": prompt}],
            temperature=0.1,
            max_tokens=max_tokens,
        )
        answer = response.choices[0].message.content
        if not answer:
            raise HTTPException(502, "The AI model returned an empty answer. Please retry.")
        return answer
    except HTTPException:
        raise
    except Exception as exc:
        msg = str(exc)
        if any(term in msg.lower() for term in ["rate_limit", "tokens per minute", "tpm"]):
            raise HTTPException(429, "Groq rate limit reached. Wait a minute and try again.")
        raise HTTPException(502, f"AI request failed: {msg[:300]}")

def read_upload(filename: str, raw: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext in {".txt", ".md", ".json", ".csv"}:
        return raw.decode("utf-8-sig", errors="replace")
    if ext == ".xlsx":
        wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"SHEET: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                values = [str(v) if v is not None else "" for v in row]
                if any(values): out.append(" | ".join(values))
        return "\n".join(out)
    if ext == ".docx":
        doc = Document(io.BytesIO(raw))
        out = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                out.append(" | ".join(c.text.replace("\n", " / ") for c in row.cells))
        return "\n".join(out)
    if ext == ".pdf":
        pdf = PdfReader(io.BytesIO(raw))
        return "\n\n".join(f"PAGE {i+1}:\n{p.extract_text() or '[No selectable text found]'}" for i, p in enumerate(pdf.pages))
    raise HTTPException(415, f"Unsupported file type: {ext}. Supported: TXT, CSV, MD, JSON, XLSX, DOCX, PDF.")

def compact_rows(rows: list[dict[str, Any]], question: str, max_chars: int = 7800):
    fields = ["line_item_id", "description", "specification", "required_unit", "required_quantity", "vendor_id", "vendor_name", "quoted", "quoted_unit", "quoted_unit_price", "currency", "discount_pct", "freight_pct", "lead_time_days", "quality_questionnaire_pass", "confidence", "notes", "source_document"]
    q = question.lower()
    if any(w in q for w in ["missing", "unquoted", "not quoted"]):
        priority = [r for r in rows if str(r.get("quoted", "")).lower() not in ("true", "1", "yes")]
    elif any(w in q for w in ["compare", "comparable", "unit", "currency", "risk", "verify", "award"]):
        priority = [r for r in rows if str(r.get("quoted_unit", "")).lower() != str(r.get("required_unit", "")).lower() or str(r.get("currency", "")).upper() not in ("INR", "") or str(r.get("confidence", "")).lower() in ("low", "medium") or bool(r.get("notes")) or str(r.get("quoted", "")).lower() not in ("true", "1", "yes")]
    else:
        priority = []
    ordered = priority + [r for r in rows if r not in priority]
    result, used = [], 0
    for row in ordered:
        item = {k: row.get(k) for k in fields if k in row}
        encoded = json.dumps(item, ensure_ascii=False, separators=(",", ":"))
        if used + len(encoded) > max_chars: continue
